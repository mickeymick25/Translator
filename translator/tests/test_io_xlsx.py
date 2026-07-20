"""
Tests for core/io_xlsx.py — XLSX I/O module.

Covers:
- clean_text(): whitespace normalization, trimming, empty strings
- load_dropdown_xlsx(): loading dropdown data, missing files, empty rows
- analyze_xlsx(): structure analysis, headers, unique values, sample data
- save_dropdown_xlsx(): multi-sheet generation, language-specific columns
"""

import pytest
from core.io_xlsx import (
    analyze_xlsx,
    clean_text,
    load_dropdown_xlsx,
    save_dropdown_xlsx,
)

# Skip all tests if openpyxl is not installed
pytestmark = pytest.mark.skipif(
    __import__("importlib").util.find_spec("openpyxl") is None,
    reason="openpyxl not installed",
)


# ─── clean_text ───────────────────────────────────────────────────


class TestCleanText:
    """Tests for clean_text()."""

    def test_trims_leading_and_trailing_whitespace(self):
        """Leading and trailing whitespace is removed."""
        assert clean_text("  hello  ") == "hello"

    def test_normalizes_internal_whitespace(self):
        """Multiple spaces between words are collapsed to one."""
        assert clean_text("hello   world") == "hello world"

    def test_normalizes_tabs_and_newlines(self):
        """Tabs and newlines are converted to single spaces."""
        assert clean_text("hello\t\nworld") == "hello world"

    def test_returns_empty_string_for_none(self):
        """None input returns an empty string."""
        assert clean_text(None) == ""

    def test_returns_empty_string_for_empty_string(self):
        """Empty string input returns an empty string."""
        assert clean_text("") == ""

    def test_returns_empty_string_for_whitespace_only(self):
        """Whitespace-only input returns an empty string."""
        assert clean_text("   \t\n  ") == ""

    def test_preserves_non_whitespace_content(self):
        """Non-whitespace content is preserved."""
        assert clean_text("Button") == "Button"

    def test_handles_numeric_input(self):
        """Numeric input is converted to string and cleaned."""
        assert clean_text(42) == "42"

    def test_handles_special_characters(self):
        """Special characters are preserved."""
        assert clean_text("100%") == "100%"
        assert clean_text("résumé") == "résumé"

    def test_mixed_whitespace_and_content(self):
        """Complex whitespace patterns are normalized."""
        assert clean_text("  hello  \t  world  \n  ") == "hello world"


# ─── load_dropdown_xlsx ───────────────────────────────────────────


class TestLoadDropdownXlsx:
    """Tests for load_dropdown_xlsx()."""

    def test_loads_valid_xlsx(self, temp_xlsx_file):
        """A valid XLSX file with headers is loaded correctly."""
        data = load_dropdown_xlsx(temp_xlsx_file)

        assert isinstance(data, list)
        assert len(data) == 4

    def test_first_entry_has_correct_fields(self, temp_xlsx_file):
        """Each entry has 'origin', 'french', and 'context' keys."""
        data = load_dropdown_xlsx(temp_xlsx_file)

        assert data[0]["origin"] == "Button"
        assert data[0]["french"] == "Bouton"
        assert data[0]["context"] == "btn"

    def test_all_entries_have_required_keys(self, temp_xlsx_file):
        """Every entry in the loaded data has the required keys."""
        data = load_dropdown_xlsx(temp_xlsx_file)

        for entry in data:
            assert "origin" in entry
            assert "french" in entry
            assert "context" in entry

    def test_contexts_are_correct(self, temp_xlsx_file):
        """Context values match the XLSX content."""
        data = load_dropdown_xlsx(temp_xlsx_file)

        contexts = {entry["context"] for entry in data}
        assert "btn" in contexts
        assert "menu" in contexts

    def test_skips_empty_origin_rows(self, temp_dir):
        """Rows with empty origin column are skipped."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Origin", "French", "Context"])
        ws.append(["Button", "Bouton", "btn"])
        ws.append(["", "", ""])  # Empty row — should be skipped
        ws.append(["File", "Fichier", "menu"])

        filepath = temp_dir / "with_empty_rows.xlsx"
        wb.save(filepath)

        data = load_dropdown_xlsx(filepath)
        assert len(data) == 2
        assert data[0]["origin"] == "Button"
        assert data[1]["origin"] == "File"

    def test_missing_context_defaults_to_unknown(self, temp_dir):
        """Rows with empty context default to 'unknown'."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Origin", "French", "Context"])
        ws.append(["Button", "Bouton", ""])  # Empty context

        filepath = temp_dir / "no_context.xlsx"
        wb.save(filepath)

        data = load_dropdown_xlsx(filepath)
        assert data[0]["context"] == "unknown"

    def test_raises_on_missing_file(self):
        """FileNotFoundError is raised when the file does not exist."""
        with pytest.raises(FileNotFoundError):
            load_dropdown_xlsx("/nonexistent/path/dropdown.xlsx")

    def test_text_is_cleaned(self, temp_dir):
        """Whitespace in cells is cleaned by clean_text()."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Origin", "French", "Context"])
        ws.append(["  Button  ", "  Bouton  ", "  btn  "])

        filepath = temp_dir / "whitespace.xlsx"
        wb.save(filepath)

        data = load_dropdown_xlsx(filepath)
        assert data[0]["origin"] == "Button"
        assert data[0]["french"] == "Bouton"
        assert data[0]["context"] == "btn"

    def test_loads_xlsx_with_string_path(self, temp_xlsx_file):
        """load_dropdown_xlsx works with a string path, not just Path."""
        data = load_dropdown_xlsx(str(temp_xlsx_file))
        assert isinstance(data, list)
        assert len(data) == 4

    def test_empty_xlsx_returns_empty_list(self, temp_dir):
        """An XLSX with only headers returns an empty list."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Origin", "French", "Context"])
        # No data rows

        filepath = temp_dir / "empty_data.xlsx"
        wb.save(filepath)

        data = load_dropdown_xlsx(filepath)
        assert data == []


# ─── analyze_xlsx ─────────────────────────────────────────────────


class TestAnalyzeXlsx:
    """Tests for analyze_xlsx()."""

    def test_returns_analysis_dict(self, temp_xlsx_file):
        """analyze_xlsx returns a dict with expected top-level keys."""
        analysis = analyze_xlsx(temp_xlsx_file)

        assert "fichier" in analysis
        assert "feuilles" in analysis
        assert "statistiques" in analysis

    def test_statistics_are_correct(self, temp_xlsx_file):
        """Statistics include total sheets, rows, and columns."""
        analysis = analyze_xlsx(temp_xlsx_file)
        stats = analysis["statistiques"]

        assert stats["total_feuilles"] == 1
        assert stats["total_lignes"] > 0
        assert stats["total_colonnes"] > 0

    def test_sheet_info_contains_required_fields(self, temp_xlsx_file):
        """Each sheet info dict has the expected fields."""
        analysis = analyze_xlsx(temp_xlsx_file)
        sheet = analysis["feuilles"][0]

        assert "nom" in sheet
        assert "dimensions" in sheet
        assert "lignes" in sheet
        assert "colonnes" in sheet
        assert "en_tetes" in sheet
        assert "nombre_lignes_donnees" in sheet
        assert "valeurs_uniques_par_colonne" in sheet

    def test_headers_are_correct(self, temp_xlsx_file):
        """Headers match the XLSX column headers."""
        analysis = analyze_xlsx(temp_xlsx_file)
        headers = analysis["feuilles"][0]["en_tetes"]

        assert headers == ["Origin", "French", "Context"]

    def test_unique_values_per_column(self, temp_xlsx_file):
        """Unique values are computed per column."""
        analysis = analyze_xlsx(temp_xlsx_file)
        unique_vals = analysis["feuilles"][0]["valeurs_uniques_par_colonne"]

        assert "Origin" in unique_vals
        assert "French" in unique_vals
        assert "Context" in unique_vals

        # Context column should have 2 unique values: btn, menu
        assert len(unique_vals["Context"]) == 2
        assert "btn" in unique_vals["Context"]
        assert "menu" in unique_vals["Context"]

    def test_data_rows_count(self, temp_xlsx_file):
        """nombre_lignes_donnees equals total rows minus header row."""
        analysis = analyze_xlsx(temp_xlsx_file)
        sheet = analysis["feuilles"][0]

        # 4 data rows in the test file
        assert sheet["nombre_lignes_donnees"] == 4

    def test_first_data_row_present(self, temp_xlsx_file):
        """First data row is included in the analysis."""
        analysis = analyze_xlsx(temp_xlsx_file)
        first_row = analysis["feuilles"][0]["premiere_ligne"]

        assert first_row is not None
        assert len(first_row) == 3  # 3 columns

    def test_raises_on_missing_file(self):
        """FileNotFoundError is raised when the file does not exist."""
        with pytest.raises(FileNotFoundError):
            analyze_xlsx("/nonexistent/path/file.xlsx")

    def test_donnees_section_present(self, temp_xlsx_file):
        """The 'donnees' section contains the full data rows."""
        analysis = analyze_xlsx(temp_xlsx_file)

        assert "donnees" in analysis
        assert len(analysis["donnees"]) > 0
        assert "feuille" in analysis["donnees"][0]
        assert "lignes" in analysis["donnees"][0]


# ─── save_dropdown_xlsx ───────────────────────────────────────────


class TestSaveDropdownXlsx:
    """Tests for save_dropdown_xlsx()."""

    def test_creates_xlsx_file(self, temp_dir, sample_dropdown_entries):
        """save_dropdown_xlsx creates an XLSX file."""
        languages = {
            "en": {
                "name": "Anglais",
                "code": "EN",
                "target": "en",
                "source_col": "origin",
            },
            "fr": {
                "name": "Français",
                "code": "FR",
                "target": "fr",
                "source_col": "french",
            },
        }
        translations = {"en": {}, "fr": {}}

        output_path = temp_dir / "output.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        assert output_path.exists()

    def test_creates_one_sheet_per_language(self, temp_dir, sample_dropdown_entries):
        """One sheet is created per language."""
        import openpyxl

        languages = {
            "en": {
                "name": "Anglais",
                "code": "EN",
                "target": "en",
                "source_col": "origin",
            },
            "fr": {
                "name": "Français",
                "code": "FR",
                "target": "fr",
                "source_col": "french",
            },
        }
        translations = {"en": {}, "fr": {}}

        output_path = temp_dir / "multi_sheet.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        wb = openpyxl.load_workbook(output_path)
        assert len(wb.sheetnames) == 2
        assert "EN" in wb.sheetnames
        assert "FR" in wb.sheetnames

    def test_en_sheet_contains_origin_values(self, temp_dir, sample_dropdown_entries):
        """The English sheet contains origin values as translations."""
        import openpyxl

        languages = {
            "en": {
                "name": "Anglais",
                "code": "EN",
                "target": "en",
                "source_col": "origin",
            },
        }
        translations = {"en": {}}

        output_path = temp_dir / "en_only.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        wb = openpyxl.load_workbook(output_path)
        ws = wb["EN"]

        # Row 1 = headers, Row 2 = first data row
        assert ws.cell(row=2, column=2).value == "Button"

    def test_fr_sheet_contains_french_values(self, temp_dir, sample_dropdown_entries):
        """The French sheet contains french column values."""
        import openpyxl

        languages = {
            "fr": {
                "name": "Français",
                "code": "FR",
                "target": "fr",
                "source_col": "french",
            },
        }
        translations = {"fr": {}}

        output_path = temp_dir / "fr_only.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        wb = openpyxl.load_workbook(output_path)
        ws = wb["FR"]

        assert ws.cell(row=2, column=2).value == "Bouton"

    def test_translated_language_sheet_uses_translations(
        self, temp_dir, sample_dropdown_entries
    ):
        """Non-EN/FR sheets use the translations dictionary."""
        import openpyxl

        languages = {
            "de": {
                "name": "Allemand",
                "code": "DE",
                "target": "de",
                "source_col": "origin",
            },
        }
        translations = {
            "de": {
                "Button": "Taste",
                "Cancel": "Abbrechen",
                "File": "Datei",
                "Open": "Öffnen",
            },
        }

        output_path = temp_dir / "de_translated.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        wb = openpyxl.load_workbook(output_path)
        ws = wb["DE"]

        assert ws.cell(row=2, column=2).value == "Taste"
        assert ws.cell(row=3, column=2).value == "Abbrechen"

    def test_falls_back_to_origin_when_translation_missing(
        self, temp_dir, sample_dropdown_entries
    ):
        """When a translation is missing, the origin text is used as fallback."""
        import openpyxl

        languages = {
            "de": {
                "name": "Allemand",
                "code": "DE",
                "target": "de",
                "source_col": "origin",
            },
        }
        translations = {
            "de": {"Button": "Taste"},  # Only Button translated
        }

        output_path = temp_dir / "partial_translation.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        wb = openpyxl.load_workbook(output_path)
        ws = wb["DE"]

        assert ws.cell(row=2, column=2).value == "Taste"
        # Cancel has no translation — falls back to origin
        assert ws.cell(row=3, column=2).value == "Cancel"

    def test_creates_parent_directories(self, temp_dir, sample_dropdown_entries):
        """save_dropdown_xlsx creates parent directories if needed."""
        languages = {
            "en": {
                "name": "Anglais",
                "code": "EN",
                "target": "en",
                "source_col": "origin",
            },
        }
        translations = {"en": {}}

        output_path = temp_dir / "nested" / "dir" / "output.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        assert output_path.exists()

    def test_headers_in_each_sheet(self, temp_dir, sample_dropdown_entries):
        """Each sheet has Origin, Traduction, Contexte as headers."""
        import openpyxl

        languages = {
            "en": {
                "name": "Anglais",
                "code": "EN",
                "target": "en",
                "source_col": "origin",
            },
        }
        translations = {"en": {}}

        output_path = temp_dir / "headers.xlsx"
        save_dropdown_xlsx(
            sample_dropdown_entries, translations, output_path, languages
        )

        wb = openpyxl.load_workbook(output_path)
        ws = wb["EN"]

        assert ws.cell(row=1, column=1).value == "Origin"
        assert ws.cell(row=1, column=2).value == "Traduction"
        assert ws.cell(row=1, column=3).value == "Contexte"


# ─── Sprint 3 - tâche 29 : garde-fou openpyxl manquant (C8) ─────────────


class TestOpenpyxlMissingGuards:
    """Vérifie que les fonctions XLSX lèvent ImportError si openpyxl est None.

    Ces tests mockent `core.io_xlsx.openpyxl` à None pour simuler un environnement
    sans openpyxl (Docker minimal, environnement réduit). Corrige C8.
    """

    def test_load_dropdown_xlsx_all_sheets_raises_importerror(self, tmp_path):
        from unittest.mock import patch

        import core.io_xlsx as iox

        # Crée un fichier XLSX valide pour que le chemin existe.
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.title = "EN"
        wb.active.append(["Origin", "Anglais", "Contexte"])
        path = tmp_path / "dummy.xlsx"
        wb.save(path)

        with patch.object(iox, "openpyxl", None):
            with pytest.raises(ImportError, match="openpyxl is required"):
                iox.load_dropdown_xlsx_all_sheets(path)

    def test_load_dropdown_xlsx_raises_importerror(self, tmp_path):
        from unittest.mock import patch

        import core.io_xlsx as iox

        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.title = "EN"
        wb.active.append(["Origin", "Anglais", "Contexte"])
        path = tmp_path / "dummy.xlsx"
        wb.save(path)

        with patch.object(iox, "openpyxl", None):
            with pytest.raises(ImportError, match="openpyxl is required"):
                iox.load_dropdown_xlsx(path)
